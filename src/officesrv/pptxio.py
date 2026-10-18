from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt

layouts = {"title": 0, "content": 1, "section": 2, "two": 3, "title_only": 5, "blank": 6}


def open_deck(path):
    if not Path(path).is_file():
        raise FileNotFoundError(f"no such file: {path}")
    try:
        return Presentation(path)
    except Exception as exc:
        raise ValueError(f"cannot open {path} as pptx: {exc}") from exc


def add_slide(deck, spec):
    kind = spec.get("layout") or ("title" if not spec.get("bullets") and not spec.get("image") and not deck.slides else "content")
    if kind not in layouts:
        raise ValueError(f"layout is one of {', '.join(layouts)}")
    slide = deck.slides.add_slide(deck.slide_layouts[layouts[kind]])
    if slide.shapes.title is not None:
        slide.shapes.title.text = spec.get("title", "")
    holders = [ph for ph in slide.placeholders if ph.placeholder_format.idx != 0]
    if spec.get("subtitle") and holders:
        holders[0].text = spec["subtitle"]
    bullets = spec.get("bullets") or []
    if bullets:
        if holders and holders[0].has_text_frame:
            frame = holders[0].text_frame
        else:
            frame = slide.shapes.add_textbox(Cm(2), Cm(4), deck.slide_width - Cm(4), Cm(10)).text_frame
        frame.clear()
        for k, item in enumerate(bullets):
            text, level = (item, 0) if isinstance(item, str) else (item.get("text", ""), item.get("level", 0))
            para = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
            para.text = text
            para.level = level
            if spec.get("font_size"):
                para.font.size = Pt(spec["font_size"])
    if spec.get("image"):
        image = Path(spec["image"])
        if not image.is_file():
            raise FileNotFoundError(f"no such image: {image}")
        width = Cm(spec["image_width_cm"]) if spec.get("image_width_cm") else None
        left = Cm(spec.get("image_left_cm", 13 if bullets else 4))
        slide.shapes.add_picture(str(image), left, Cm(spec.get("image_top_cm", 4.5)), width=width)
    if spec.get("table"):
        grid = spec["table"]
        ncols = max(len(r) for r in grid)
        frame = slide.shapes.add_table(len(grid), ncols, Cm(2), Cm(4.5), deck.slide_width - Cm(4), Cm(1) * len(grid))
        for r, row in enumerate(grid):
            for c in range(ncols):
                frame.table.cell(r, c).text = "" if c >= len(row) else str(row[c])
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]
    return slide


def deck_text(deck, max_chars):
    chunks = []
    for n, slide in enumerate(deck.slides, start=1):
        chunks.append(f"## slide {n}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text)
            elif shape.has_table:
                chunks.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
            elif shape.shape_type == 13:
                chunks.append("[image]")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            chunks.append("notes: " + slide.notes_slide.notes_text_frame.text)
    text = "\n".join(chunks)
    return text if len(text) <= max_chars else text[:max_chars] + f"\n[cut at {max_chars} of {len(text)} chars]"
