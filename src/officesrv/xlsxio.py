from datetime import date, datetime, time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

thin = Side(style="thin")


def open_book(path, values_only=False):
    try:
        return load_workbook(path, data_only=values_only)
    except Exception as exc:
        raise ValueError(f"cannot open {path} as xlsx: {exc}") from exc


def sheet_of(book, name):
    if name not in book.sheetnames:
        raise ValueError(f"no sheet {name!r}, sheets are {book.sheetnames}")
    return book[name]


def write_rows(ws, grid, start_row=1):
    for r, row in enumerate(grid, start=start_row):
        for c, val in enumerate(row, start=1):
            ws.cell(r, c, val)


def plain(val):
    if isinstance(val, (datetime, date, time)):
        return val.isoformat()
    return val


def read_rows(ws, max_cells):
    out, seen = [], 0
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        while cells and cells[-1] is None:
            cells.pop()
        seen += max(1, len(cells))
        if seen > max_cells:
            break
        out.append([plain(v) for v in cells])
    while out and not out[-1]:
        out.pop()
    return out


def style_header(ws, nrows=1, fill="4472C4"):
    for r in range(1, nrows + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def grid_borders(ws):
    edge = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = edge


def fit_widths(ws):
    for c in range(1, ws.max_column + 1):
        letter = get_column_letter(c)
        longest = max((len(str(cell.value)) for cell in ws[letter] if cell.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(longest + 2, 8), 60)


def new_book():
    book = Workbook()
    book.remove(book.active)
    return book
