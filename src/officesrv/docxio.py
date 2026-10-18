from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

aligns = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
          "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def open_doc(path):
    try:
        return Document(path)
    except Exception as exc:
        raise ValueError(f"cannot open {path} as docx: {exc}") from exc


def add_para(doc, spec):
    if not isinstance(spec, dict):
        spec = {"text": str(spec)}
    style = spec.get("style")
    if style and style.lower().startswith("heading"):
        level = int(style[-1]) if style[-1].isdigit() else 1
        p = doc.add_heading(spec.get("text", ""), level=level)
    else:
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        piece = p.add_run(spec.get("text", ""))
        if spec.get("size"):
            piece.font.size = Pt(spec["size"])
        piece.bold = bool(spec.get("bold"))
        piece.italic = bool(spec.get("italic"))
    if spec.get("align"):
        if spec["align"] not in aligns:
            raise ValueError(f"align is one of {', '.join(aligns)}")
        p.alignment = aligns[spec["align"]]
    return p


def add_table(doc, grid, header=True, widths=None):
    if not grid or not any(grid):
        raise ValueError("table needs at least one non empty row")
    ncols = max(len(r) for r in grid)
    tbl = doc.add_table(rows=len(grid), cols=ncols)
    tbl.style = "Table Grid"
    for r, row in enumerate(grid):
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.text = "" if c >= len(row) or row[c] is None else str(row[c])
            if header and r == 0:
                for piece in cell.paragraphs[0].runs:
                    piece.bold = True
    if widths:
        for c, w in enumerate(widths[:ncols]):
            for row in tbl.rows:
                row.cells[c].width = Cm(w)
    return tbl


def doc_text(doc, max_chars):
    chunks = []
    body = doc.element.body
    tables = iter(doc.tables)
    paras = iter(doc.paragraphs)
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = next(paras)
            if p.text.strip():
                chunks.append(("# " if p.style.name.startswith("Heading") else "") + p.text)
        elif tag == "tbl":
            for row in next(tables).rows:
                chunks.append(" | ".join(c.text for c in row.cells))
    text = "\n".join(chunks)
    return text if len(text) <= max_chars else text[:max_chars] + f"\n[cut at {max_chars} of {len(text)} chars]"


def all_paragraphs(doc):
    yield from doc.paragraphs
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_text(doc, old, new):
    if not old:
        raise ValueError("text to replace is empty")
    hits = 0
    for p in all_paragraphs(doc):
        if old not in p.text:
            continue
        count = p.text.count(old)
        if any(old in piece.text for piece in p.runs):
            for piece in p.runs:
                piece.text = piece.text.replace(old, new)
        else:
            # match spans several runs, keep the first run formatting
            merged = p.text.replace(old, new)
            for piece in p.runs[1:]:
                piece.text = ""
            p.runs[0].text = merged
        hits += count
    return hits
